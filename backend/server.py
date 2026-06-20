import os
import sys
import json
import sqlite3
import random
from datetime import datetime
from flask import Flask, request, jsonify
from flask_cors import CORS

# Add current folder to path
sys.path.append(os.path.dirname(os.path.abspath(__file__)))

app = Flask(__name__)
CORS(app) # Enable CORS for frontend connection

# Import rule engine from main_standalone
try:
    from main_standalone import audit_single_invoice as run_rule_checks, HSN_MASTER_REGISTRY
except ImportError:
    # Inline fallback if main_standalone is missing
    HSN_MASTER_REGISTRY = {
        "2523": {"expected_rate": "28%", "category": "Cement"},
        "1512": {"expected_rate": "5%",  "category": "Sunflower Oil"},
        "1516": {"expected_rate": "12%", "category": "Vegetable Fats/Hydrogenated Oils"},
        "8471": {"expected_rate": "18%", "category": "Computers/Electronics"},
        "9965": {"expected_rate": "5%",  "category": "Goods Transport Services (GTA)"},
    }
    def run_rule_checks(invoice_data):
        return []


# Import VLM OCR pipeline with graceful fallback
has_vlm = False
try:
    import gst_ocr_pipeline as pipeline
    has_vlm = True
    DB_PATH = pipeline.DB_PATH
except Exception as e:
    print(f"[WARNING] VLM OCR Pipeline import failed: {e}. Fallback mode active.")
    pipeline = None
    DB_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), "gst_purchase_registry.db")

# SQLite fallback helpers if pipeline is not imported
def fallback_init_db():
    conn = sqlite3.connect(DB_PATH)
    conn.execute("PRAGMA foreign_keys = ON")
    conn.execute("""
        CREATE TABLE IF NOT EXISTS purchase_registry (
            id                    INTEGER PRIMARY KEY AUTOINCREMENT,
            gstin_of_supplier     TEXT NOT NULL,
            trade_legal_name      TEXT,
            type_of_inward_supply TEXT,
            document_type         TEXT,
            document_number       TEXT NOT NULL,
            document_date         TEXT,
            taxable_value         REAL DEFAULT 0,
            integrated_tax        REAL DEFAULT 0,
            central_tax           REAL DEFAULT 0,
            state_ut_tax          REAL DEFAULT 0,
            cess                  REAL DEFAULT 0,
            buyer_gstin           TEXT,
            place_of_supply       TEXT,
            grand_total           REAL DEFAULT 0,
            confidence_score      REAL,
            source_file           TEXT,
            return_period         TEXT,
            needs_review          INTEGER DEFAULT 0,
            raw_json              TEXT,
            created_at            TEXT DEFAULT (datetime('now')),
            UNIQUE(gstin_of_supplier, document_number, document_type)
        )
    """)
    conn.execute("""
        CREATE TABLE IF NOT EXISTS purchase_registry_items (
            id                    INTEGER PRIMARY KEY AUTOINCREMENT,
            purchase_registry_id  INTEGER NOT NULL REFERENCES purchase_registry(id) ON DELETE CASCADE,
            sr_no                 TEXT,
            description           TEXT,
            hsn_code              TEXT,
            quantity              TEXT,
            unit                  TEXT,
            rate                  TEXT,
            taxable_amount        TEXT,
            gst_rate              TEXT,
            gst_amount            TEXT,
            total                 TEXT
        )
    """)
    conn.commit()
    conn.close()

def fallback_save_to_purchase_registry(record):
    fallback_init_db()
    conn = sqlite3.connect(DB_PATH)
    try:
        confidence = (record.get("confidence") or {}).get("overall", 0.95)
        warnings_list = record.get("warnings", [])
        needs_review = 1 if (warnings_list or confidence < 0.7) else 0

        def to_float(v):
            try: return float(v)
            except: return 0.0

        params = {
            "gstin_of_supplier": record.get("gstin_of_supplier", ""),
            "trade_legal_name": record.get("trade_legal_name", ""),
            "type_of_inward_supply": record.get("type_of_inward_supply", "Inputs"),
            "document_type": record.get("document_type", "Invoice"),
            "document_number": record.get("document_number", ""),
            "document_date": record.get("document_date", ""),
            "taxable_value": to_float(record.get("taxable_value")),
            "integrated_tax": to_float(record.get("integrated_tax")),
            "central_tax": to_float(record.get("central_tax")),
            "state_ut_tax": to_float(record.get("state_ut_tax")),
            "cess": to_float(record.get("cess")),
            "buyer_gstin": record.get("buyer_gstin", ""),
            "place_of_supply": record.get("place_of_supply", ""),
            "grand_total": to_float(record.get("grand_total")),
            "confidence_score": confidence,
            "source_file": record.get("source_file", ""),
            "return_period": record.get("return_period", ""),
            "needs_review": needs_review,
            "raw_json": json.dumps(record),
        }

        # Check existing
        existing = conn.execute(
            "SELECT id FROM purchase_registry WHERE gstin_of_supplier=:gstin_of_supplier "
            "AND document_number=:document_number AND document_type=:document_type",
            params,
        ).fetchone()

        if existing:
            registry_id = existing[0]
            params["id"] = registry_id
            conn.execute("""
                UPDATE purchase_registry SET
                    trade_legal_name=:trade_legal_name,
                    type_of_inward_supply=:type_of_inward_supply,
                    document_date=:document_date,
                    taxable_value=:taxable_value,
                    integrated_tax=:integrated_tax,
                    central_tax=:central_tax,
                    state_ut_tax=:state_ut_tax,
                    cess=:cess,
                    buyer_gstin=:buyer_gstin,
                    place_of_supply=:place_of_supply,
                    grand_total=:grand_total,
                    confidence_score=:confidence_score,
                    source_file=:source_file,
                    return_period=:return_period,
                    needs_review=:needs_review,
                    raw_json=:raw_json
                WHERE id=:id
            """, params)
            conn.execute("DELETE FROM purchase_registry_items WHERE purchase_registry_id=?", (registry_id,))
        else:
            cur = conn.execute("""
                INSERT INTO purchase_registry (
                    gstin_of_supplier, trade_legal_name, type_of_inward_supply, document_type,
                    document_number, document_date, taxable_value, integrated_tax, central_tax,
                    state_ut_tax, cess, buyer_gstin, place_of_supply, grand_total,
                    confidence_score, source_file, return_period, needs_review, raw_json
                ) VALUES (
                    :gstin_of_supplier, :trade_legal_name, :type_of_inward_supply, :document_type,
                    :document_number, :document_date, :taxable_value, :integrated_tax, :central_tax,
                    :state_ut_tax, :cess, :buyer_gstin, :place_of_supply, :grand_total,
                    :confidence_score, :source_file, :return_period, :needs_review, :raw_json
                )
            """, params)
            registry_id = cur.lastrowid

        for item in record.get("line_items", []):
            conn.execute("""
                INSERT INTO purchase_registry_items (
                    purchase_registry_id, sr_no, description, hsn_code, quantity, unit,
                    rate, taxable_amount, gst_rate, gst_amount, total
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?)
            """, (
                registry_id,
                item.get("sr", ""), item.get("description", ""), item.get("hsn_code", ""),
                item.get("quantity", ""), item.get("unit", ""), item.get("rate", ""),
                item.get("taxable_amount", ""), item.get("gst_rate", ""),
                item.get("gst_amount", ""), item.get("total", ""),
            ))
        conn.commit()
        return registry_id
    finally:
        conn.close()

def fallback_get_purchase_registry():
    fallback_init_db()
    conn = sqlite3.connect(DB_PATH)
    conn.row_factory = sqlite3.Row
    try:
        rows = conn.execute("SELECT * FROM purchase_registry ORDER BY id DESC").fetchall()
        results = []
        for row in rows:
            rec = dict(row)
            items = conn.execute(
                "SELECT sr_no AS sr, description, hsn_code, quantity, unit, rate, "
                "taxable_amount, gst_rate, gst_amount, total "
                "FROM purchase_registry_items WHERE purchase_registry_id = ?",
                (row["id"],),
            ).fetchall()
            rec["line_items"] = [dict(i) for i in items]
            try:
                raw = json.loads(rec.get("raw_json") or "{}")
                rec["warnings"] = raw.get("warnings", [])
            except:
                rec["warnings"] = []
            results.append(rec)
        return results
    finally:
        conn.close()

# Initial database setup
if pipeline:
    pipeline.init_db()
else:
    fallback_init_db()


# --- ROUTES ---

@app.route('/api/health', methods=['GET'])
def health():
    ollama_active = False
    try:
        from openai import OpenAI
        client = OpenAI(base_url="http://localhost:11434/v1", api_key="ollama")
        client.models.list()
        ollama_active = True
    except Exception:
        pass

    return jsonify({
        "status": "healthy",
        "vlm_loaded": has_vlm,
        "ollama_active": ollama_active,
        "database_path": str(DB_PATH),
        "mode": "VLM-premium" if has_vlm else "VLM-fallback-mode"
    })

@app.route('/api/upload-invoice', methods=['POST'])
def upload_invoice():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400

    filename = file.filename.lower()
    ocr_mode = request.form.get('ocr_mode', 'auto')
    
    # Save file directly into the invoices directory
    invoices_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "invoices")
    os.makedirs(invoices_dir, exist_ok=True)
    saved_path = os.path.join(invoices_dir, file.filename)
    file.save(saved_path)

    record = None
    vlm_warning = None
    
    try:
        if ocr_mode == 'vlm' or (ocr_mode == 'auto' and has_vlm):
            if has_vlm and pipeline:
                try:
                    buyer_gstin = request.form.get('buyer_gstin', '')
                    record = pipeline.process_invoice_to_registry(
                        saved_path,
                        filename_hint=file.filename,
                        save_to_db=True,
                        db_path=DB_PATH,
                        buyer_gstin=buyer_gstin
                    )
                except Exception as vlm_err:
                    print(f"[VLM ERROR] Qwen local model failed: {vlm_err}")
                    vlm_warning = f"Qwen2.5-VL neural model failed to initialize/load: {str(vlm_err)}."
            else:
                vlm_warning = "Qwen2.5-VL OCR pipeline is not loaded on this backend."

        if record is None:
            return jsonify({"error": vlm_warning or "VLM OCR extraction failed. Please check backend logs."}), 500

        # Execute rules audit checks
        rule_flags = run_rule_checks(record)
        existing_warnings = record.get("warnings", [])
        if not isinstance(existing_warnings, list):
            existing_warnings = [existing_warnings] if existing_warnings else []
            
        audit_warnings = [f["description"] for f in rule_flags] if rule_flags else []
        record["warnings"] = list(set(existing_warnings + audit_warnings))
        
        if vlm_warning:
            record["warnings"].append(vlm_warning)

        # Write output JSON to processed_invoices folder
        processed_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "processed_invoices")
        os.makedirs(processed_dir, exist_ok=True)
        file_stem = os.path.splitext(file.filename)[0]
        json_path = os.path.join(processed_dir, f"{file_stem}.json")
        with open(json_path, "w", encoding="utf-8") as f:
            json.dump(record, f, ensure_ascii=False, indent=2)

        # Update combined purchase_registry.json
        combined_path = os.path.join(processed_dir, "purchase_registry.json")
        try:
            if pipeline:
                all_records = pipeline.get_purchase_registry(db_path=DB_PATH)
            else:
                all_records = fallback_get_purchase_registry()
            with open(combined_path, "w", encoding="utf-8") as f:
                json.dump(all_records, f, ensure_ascii=False, indent=2)
        except Exception as comb_err:
            print(f"[WARNING] Failed to update combined purchase_registry.json: {comb_err}")
        
        return jsonify(record)
    except Exception as err:
        return jsonify({"error": f"OCR pipeline failed: {str(err)}"}), 500

@app.route('/api/parse-gstr2b', methods=['POST'])
def parse_gstr2b():
    if 'file' not in request.files:
        return jsonify({"error": "No file uploaded"}), 400
    
    file = request.files['file']
    if file.filename == '':
        return jsonify({"error": "No file selected"}), 400

    temp_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "temp")
    os.makedirs(temp_dir, exist_ok=True)
    temp_path = os.path.join(temp_dir, file.filename)
    file.save(temp_path)

    try:
        # Check standard xlsx parsing
        if pipeline:
            results = pipeline.parse_gstr2b_excel(temp_path)
        else:
            # Fallback simple Excel reader using openpyxl
            import openpyxl
            wb = openpyxl.load_workbook(temp_path, data_only=True)
            results = []
            for sheet_name in ("B2B", "b2b", "B2B Invoice Details", wb.sheetnames[0]):
                if sheet_name not in wb.sheetnames:
                    continue
                ws = wb[sheet_name]
                rows = list(ws.iter_rows(values_only=True))
                if len(rows) < 2:
                    continue
                
                # Check for header row
                header_idx = 0
                for idx, row in enumerate(rows[:5]):
                    if any(row and any(isinstance(val, str) and "invoice" in val.lower() for val in row if val)):
                        header_idx = idx
                        break
                
                headers = [str(h).strip() if h else "" for h in rows[header_idx]]
                
                def col_val(row, label):
                    try:
                        idx = next(i for i, h in enumerate(headers) if label.lower() in h.lower())
                        val = row[idx]
                        return str(val).strip() if val is not None else ""
                    except StopIteration:
                        return ""

                for row in rows[header_idx+1:]:
                    if not any(row):
                        continue
                    
                    inv_num = col_val(row, "Invoice Number") or col_val(row, "Invoice No") or col_val(row, "Document Number")
                    if not inv_num:
                        continue

                    inv = {
                        "source_file": file.filename,
                        "doc_type": col_val(row, "Document Type") or "TAX_INVOICE",
                        "supplier_gstin": col_val(row, "GSTIN of Supplier") or col_val(row, "GSTIN"),
                        "supplier_name": col_val(row, "Trade/Legal Name") or col_val(row, "Supplier Name"),
                        "invoice_number": inv_num,
                        "invoice_date": col_val(row, "Invoice Date") or col_val(row, "Document Date"),
                        "place_of_supply": col_val(row, "Place of Supply"),
                        "grand_total": col_val(row, "Invoice Value") or col_val(row, "Grand Total"),
                        "cgst_amount": col_val(row, "Central Tax") or col_val(row, "CGST"),
                        "sgst_amount": col_val(row, "State/UT Tax") or col_val(row, "SGST"),
                        "igst_amount": col_val(row, "Integrated Tax") or col_val(row, "IGST"),
                        "cess_amount": col_val(row, "Cess"),
                        "taxable_value": col_val(row, "Taxable Value"),
                        "line_items": [],
                        "confidence": {"overall": 1.0}
                    }
                    results.append(inv)
                break
        
        if os.path.exists(temp_path):
            os.remove(temp_path)

        return jsonify(results)
    except Exception as err:
        if os.path.exists(temp_path):
            os.remove(temp_path)
        return jsonify({"error": f"Failed parsing GSTR-2B Excel: {str(err)}"}), 500

@app.route('/api/purchase-registry', methods=['GET'])
def get_registry():
    try:
        if pipeline:
            records = pipeline.get_purchase_registry(db_path=DB_PATH)
        else:
            records = fallback_get_purchase_registry()
        return jsonify(records)
    except Exception as err:
        return jsonify({"error": f"Database read error: {str(err)}"}), 500

@app.route('/api/purchase-registry', methods=['POST'])
def create_purchase_record():
    data = request.json
    if not data:
        return jsonify({"error": "No data provided"}), 400
        
    trade_legal_name = data.get("supplierName") or data.get("trade_legal_name", "")
    gstin_of_supplier = data.get("supplierGstin") or data.get("gstin_of_supplier", "")
    document_number = data.get("invoiceNumber") or data.get("document_number", "")
    document_date = data.get("invoiceDate") or data.get("document_date", "")
    
    if not gstin_of_supplier or not document_number:
        return jsonify({"error": "Supplier GSTIN and Invoice Number are required"}), 400
        
    def to_float(v):
        try: return float(v)
        except: return 0.0
        
    taxable_value = to_float(data.get("taxableValue") or data.get("taxable_value", 0))
    central_tax = to_float(data.get("cgst") or data.get("central_tax", 0))
    state_ut_tax = to_float(data.get("sgst") or data.get("state_ut_tax", 0))
    integrated_tax = to_float(data.get("igst") or data.get("integrated_tax", 0))
    cess = to_float(data.get("cess", 0))
    grand_total = to_float(data.get("totalAmount") or data.get("grand_total", 0))
    
    # Run audit rules
    audit_record = {
        "supplier_gstin": gstin_of_supplier,
        "buyer_gstin": data.get("buyer_gstin", "09AAAAC4451M1Z1"),
        "taxable_value": taxable_value,
        "central_tax": central_tax,
        "state_ut_tax": state_ut_tax,
        "integrated_tax": integrated_tax,
        "grand_total": grand_total
    }
    
    warnings_list = []
    try:
        rule_flags = run_rule_checks(audit_record)
        warnings_list = [f["description"] for f in rule_flags] if rule_flags else []
    except Exception as e:
        print(f"[WARNING] Run rules failed: {e}")

    needs_review = 1 if warnings_list else 0
    
    # Derive return period
    return_period = ""
    if document_date:
        for fmt in ("%d-%b-%Y", "%d-%m-%Y", "%Y-%m-%d"):
            try:
                return_period = datetime.strptime(document_date.strip(), fmt).strftime("%m-%Y")
                break
            except ValueError:
                continue
    if not return_period:
        return_period = datetime.now().strftime("%m-%Y")

    record = {
        "buyer_gstin": data.get("buyer_gstin", "09AAAAC4451M1Z1"),
        "gstin_of_supplier": gstin_of_supplier,
        "trade_legal_name": trade_legal_name,
        "type_of_inward_supply": "Inputs",
        "document_type": "Invoice",
        "document_number": document_number,
        "document_date": document_date,
        "taxable_value": taxable_value,
        "integrated_tax": integrated_tax,
        "central_tax": central_tax,
        "state_ut_tax": state_ut_tax,
        "cess": cess,
        "grand_total": grand_total,
        "confidence_score": 1.0,
        "source_file": "Manual Entry",
        "return_period": return_period,
        "needs_review": needs_review,
        "warnings": warnings_list,
        "line_items": data.get("line_items") or [
            {
                "sr": "1",
                "description": "Goods / Supplies",
                "hsn_code": data.get("hsnCode", "0000"),
                "taxable_amount": str(taxable_value),
                "gst_rate": str(data.get("gstRate", 18)) + "%",
                "gst_amount": str(central_tax + state_ut_tax + integrated_tax),
                "total": str(grand_total)
            }
        ]
    }
    
    try:
        if pipeline:
            record_id = pipeline.save_to_purchase_registry(record, db_path=DB_PATH)
        else:
            record_id = fallback_save_to_purchase_registry(record)
            
        record["id"] = record_id
        
        # After inserting a manual invoice, we should sync/update purchase_registry.json flat file
        try:
            processed_dir = os.path.join(os.path.dirname(os.path.abspath(__file__)), "processed_invoices")
            os.makedirs(processed_dir, exist_ok=True)
            combined_path = os.path.join(processed_dir, "purchase_registry.json")
            if pipeline:
                all_records = pipeline.get_purchase_registry(db_path=DB_PATH)
            else:
                all_records = fallback_get_purchase_registry()
            with open(combined_path, "w", encoding="utf-8") as f:
                json.dump(all_records, f, ensure_ascii=False, indent=2)
        except Exception as comb_err:
            print(f"[WARNING] Failed to update combined purchase_registry.json: {comb_err}")

        return jsonify(record)
    except Exception as err:
        return jsonify({"error": f"Failed to save manual record: {str(err)}"}), 500

@app.route('/api/purchase-registry/<int:record_id>', methods=['PUT'])
def update_purchase_record(record_id):
    data = request.json
    if not data:
        return jsonify({"error": "No data provided"}), 400
    
    conn = sqlite3.connect(DB_PATH)
    try:
        existing = conn.execute("SELECT id FROM purchase_registry WHERE id = ?", (record_id,)).fetchone()
        if not existing:
            return jsonify({"error": f"Record with ID {record_id} not found"}), 404
        
        trade_legal_name = data.get("supplierName") or data.get("trade_legal_name", "")
        gstin_of_supplier = data.get("supplierGstin") or data.get("gstin_of_supplier", "")
        document_number = data.get("invoiceNumber") or data.get("document_number", "")
        document_date = data.get("invoiceDate") or data.get("document_date", "")
        
        def to_float(v):
            try: return float(v)
            except: return 0.0
            
        taxable_value = to_float(data.get("taxableValue") or data.get("taxable_value", 0))
        central_tax = to_float(data.get("cgst") or data.get("central_tax", 0))
        state_ut_tax = to_float(data.get("sgst") or data.get("state_ut_tax", 0))
        integrated_tax = to_float(data.get("igst") or data.get("integrated_tax", 0))
        cess = to_float(data.get("cess", 0))
        grand_total = to_float(data.get("totalAmount") or data.get("grand_total", 0))
        
        audit_record = {
            "supplier_gstin": gstin_of_supplier,
            "buyer_gstin": data.get("buyer_gstin", "09AAAAC4451M1Z1"),
            "taxable_value": taxable_value,
            "central_tax": central_tax,
            "state_ut_tax": state_ut_tax,
            "integrated_tax": integrated_tax,
            "grand_total": grand_total
        }
        rule_flags = run_rule_checks(audit_record)
        warnings_list = [f["description"] for f in rule_flags] if rule_flags else []
        needs_review = 1 if (warnings_list or to_float(data.get("confidence_score", 1.0)) < 0.7) else 0
        buyer_gstin = data.get("buyer_gstin", "09AAAAC4451M1Z1")
        
        conn.execute("""
            UPDATE purchase_registry SET
                trade_legal_name = ?,
                gstin_of_supplier = ?,
                document_number = ?,
                document_date = ?,
                taxable_value = ?,
                central_tax = ?,
                state_ut_tax = ?,
                integrated_tax = ?,
                cess = ?,
                grand_total = ?,
                needs_review = ?,
                buyer_gstin = ?,
                raw_json = ?
            WHERE id = ?
        """, (
            trade_legal_name,
            gstin_of_supplier,
            document_number,
            document_date,
            taxable_value,
            central_tax,
            state_ut_tax,
            integrated_tax,
            cess,
            grand_total,
            needs_review,
            buyer_gstin,
            json.dumps(data),
            record_id
        ))
        
        line_items = data.get("line_items", [])
        if line_items:
            conn.execute("DELETE FROM purchase_registry_items WHERE purchase_registry_id = ?", (record_id,))
            for item in line_items:
                conn.execute("""
                    INSERT INTO purchase_registry_items (
                        purchase_registry_id, sr_no, description, hsn_code, quantity, unit,
                        rate, taxable_amount, gst_rate, gst_amount, total
                    ) VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """, (
                    record_id,
                    item.get("sr") or item.get("sr_no", "1"),
                    item.get("description", ""),
                    item.get("hsn_code", ""),
                    item.get("quantity", ""),
                    item.get("unit", ""),
                    item.get("rate", ""),
                    item.get("taxable_amount", ""),
                    item.get("gst_rate", ""),
                    item.get("gst_amount", ""),
                    item.get("total", ""),
                ))
                
        conn.commit()
        
        updated_record = {
            "id": record_id,
            "trade_legal_name": trade_legal_name,
            "gstin_of_supplier": gstin_of_supplier,
            "document_number": document_number,
            "document_date": document_date,
            "taxable_value": taxable_value,
            "central_tax": central_tax,
            "state_ut_tax": state_ut_tax,
            "integrated_tax": integrated_tax,
            "cess": cess,
            "grand_total": grand_total,
            "warnings": warnings_list,
            "needs_review": needs_review,
            "line_items": line_items
        }
        return jsonify(updated_record)
    except Exception as err:
        conn.rollback()
        return jsonify({"error": f"Failed to update database: {str(err)}"}), 500
    finally:
        conn.close()

@app.route('/api/purchase-registry/<int:record_id>', methods=['DELETE'])
def delete_purchase_record(record_id):
    conn = sqlite3.connect(DB_PATH)
    try:
        existing = conn.execute("SELECT id FROM purchase_registry WHERE id = ?", (record_id,)).fetchone()
        if not existing:
            return jsonify({"error": f"Record with ID {record_id} not found"}), 404
        
        conn.execute("PRAGMA foreign_keys = ON")
        conn.execute("DELETE FROM purchase_registry WHERE id = ?", (record_id,))
        conn.commit()
        return jsonify({"message": f"Record {record_id} deleted successfully", "id": record_id})
    except Exception as err:
        conn.rollback()
        return jsonify({"error": f"Failed to delete record from database: {str(err)}"}), 500
    finally:
        conn.close()

@app.route('/api/run-rules', methods=['POST'])
def run_rules():
    invoice_data = request.json
    if not invoice_data:
        return jsonify({"error": "No invoice data provided"}), 400
    try:
        flags = run_rule_checks(invoice_data)
        return jsonify(flags)
    except Exception as err:
        return jsonify({"error": f"Rule check error: {str(err)}"}), 500

@app.route('/api/chat', methods=['POST'])
def chat():
    data = request.json or {}
    message = data.get("message", "")
    history = data.get("history", [])
    context = data.get("context", {})

    if not message:
        return jsonify({"error": "Message is required"}), 400

    system_instruction = (
        "You are an interactive AI GST Assistant for Indian business owners.\n"
        "You have access to the invoice reconciliation context provided below.\n\n"
        "STRICT LANGUAGE RULE:\n"
        "1. Look at the language the user just used in their latest message.\n"
        "2. If the user writes in English, you MUST reply ONLY in clear English.\n"
        "3. If the user writes in Hinglish/Hindi, you MUST reply in Hinglish/Hindi.\n"
        "4. Never mix multiple languages or hallucinate foreign laws in a response.\n\n"
        "Explain tax details, HSN discrepancies, and supplier default status using plain terms. Keep your answers brief and simple."
    )

    try:
        from openai import OpenAI
        client = OpenAI(base_url="http://localhost:11434/v1", api_key="ollama")
        
        # Build chat payload for Ollama
        chat_history = [{"role": "system", "content": system_instruction}]
        
        # Include context for Ollama to evaluate
        if context:
            context_str = f"Current context: {json.dumps(context)}. "
            chat_history.append({"role": "system", "content": context_str})
            
        for h in history:
            chat_history.append({"role": h.get("role"), "content": h.get("content")})
            
        chat_history.append({"role": "user", "content": message})
        
        response = client.chat.completions.create(
            model="gemma:7b",
            messages=chat_history
        )
        ai_reply = response.choices[0].message.content
        return jsonify({"reply": ai_reply, "source": "ollama"})
        
    except Exception as err:
        # Graceful sandbox chat reply generator
        msg_lower = message.lower()
        reply_lang = "en"
        if any(w in msg_lower for w in ["hai", "kya", "bhai", "karo", "theek", "mismatch"]):
            reply_lang = "hing"
        elif any(w in msg_lower for w in ["क्या", "कैसे", "नमस्ते", "जीएसटी"]):
            reply_lang = "hi"
            
        # Context extraction details for smart fallback chat
        mismatch_summary = ""
        if context and isinstance(context, list):
            mismatch_count = len([i for i in context if i.get("status") != "MATCHED"])
            mismatch_summary = f" I see {mismatch_count} discrepancies in your reconciliation."

        if reply_lang == "hing":
            reply = (
                f"Namaste! Mujhe lagta hai aap compliance ke bare me puch rahe hain.{mismatch_summary}\n\n"
                "Aapke invoices me HSN code, state mismatch, ya missing supplier upload ke rules check kiye gaye hain. "
                "Mujhse koi bhi details puchiye.\n\n"
                "*(Note: Local Assistant is offline. Active sandbox assistant mode is running. Start Assistant to connect full chatbot)*"
            )
        elif reply_lang == "hi":
            reply = (
                f"नमस्ते! जीएसटी नियमों को समझने में मैं आपकी मदद कर सकता हूँ।{mismatch_summary}\n\n"
                "मैंने आपके बिलों की जांच की है: एचएसएन कोड अंतर और कर राशि विसंगति को हल करने के लिए आप सप्लायर से संपर्क कर सकते हैं।\n\n"
                "*(नोट: लोकल Assistant ऑफलाइन है। सैंडबॉक्स असिस्टेंट सक्रिय है। कृपया बैकग्राउंड में Assistant चलाएं)*"
            )
        else:
            reply = (
                f"Hello! I am your AI GST Tax assistant. I can help analyze your discrepancies.{mismatch_summary}\n\n"
                "Common issues like HSN mismatch (e.g., milk HSN 0401 vs service code 9987) or Supplier defaults "
                "should be checked using our Action Center. You can send supplier notifications straight over WhatsApp.\n\n"
                "*(Note: Local Assistant is offline. Running in Sandbox AI mode. Start Assistant to activate full AI models)*"
            )
            
        return jsonify({"reply": reply, "source": "sandbox"})

if __name__ == "__main__":
    port = int(os.environ.get("PORT", 5000))
    print(f"[INFO] PocketCA Backend REST Server running on port {port}...")
    app.run(host="127.0.0.1", port=port, debug=True)
